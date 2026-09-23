"""Bounded tabular adapters for check-in summary imports.

Every supported input is normalized here before identity resolution or database
writes.  Adapters never guess at unknown columns and never evaluate workbook
formulas.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import PurePosixPath
from typing import Literal
from urllib.parse import parse_qs, urlencode, urlsplit, urlunsplit
from zipfile import BadZipFile, ZipFile

import httpx
from openpyxl import load_workbook

_MAX_TEXT_BYTES = 2 * 1024 * 1024
_MAX_XLSX_BYTES = 5 * 1024 * 1024
_MAX_GOOGLE_BYTES = _MAX_TEXT_BYTES
_MAX_ROWS = 10_000
_MAX_COLUMNS = 20
_MAX_HEADER_LENGTH = 128
_MAX_CELL_LENGTH = 512
_MAX_XLSX_ENTRIES = 1_000
_MAX_XLSX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
_MAX_XLSX_COMPRESSION_RATIO = 100
_GOOGLE_HOST = "docs.google.com"
_GOOGLE_PATH = re.compile(r"^/spreadsheets/d/([A-Za-z0-9_-]{6,128})(?:/[^?]*)?$")
_GOOGLE_EXPORT_HOST = re.compile(r"^doc-[a-z0-9-]{1,96}-sheets\.googleusercontent\.com$")
_TWITCH_LOGIN = re.compile(r"^[A-Za-z0-9_]{1,25}$")


class CheckinImportValidationError(ValueError):
    """The supplied table cannot be imported without guessing or unsafe work."""


@dataclass(frozen=True, slots=True)
class SummaryRow:
    source_row: int
    platform_user_id: str | None
    username: str | None
    display_name: str | None
    total_days: int
    last_checkin_date: date
    current_streak: int | None
    daily_order: int | None

    @property
    def source_key(self) -> str:
        if self.platform_user_id:
            return f"id:{self.platform_user_id}"
        if self.username:
            return f"username:{self.username.casefold()}"
        return f"display:{(self.display_name or '').casefold()}"


@dataclass(frozen=True, slots=True)
class InvalidSummaryRow:
    source_row: int
    issue: str


@dataclass(frozen=True, slots=True)
class ParsedSummary:
    format: Literal["csv", "tsv", "xlsx", "google_sheets"]
    sheet_name: str | None
    content_sha256: str
    column_mapping: tuple[tuple[str, int], ...]
    rows: tuple[SummaryRow | InvalidSummaryRow, ...]


@dataclass(frozen=True, slots=True)
class InspectedSummary:
    format: Literal["csv", "tsv", "xlsx", "google_sheets"]
    sheet_name: str | None
    content_sha256: str
    headers: tuple[str, ...]
    suggested_mapping: dict[str, int]


_HEADER_ALIASES: dict[str, frozenset[str]] = {
    "platform_user_id": frozenset(
        {
            "platformuserid",
            "userid",
            "twitchid",
            "twitchuserid",
            "使用者id",
            "用戶id",
        }
    ),
    "username": frozenset(
        {"username", "user", "username/login", "login", "twitchusername", "使用者名稱", "帳號"}
    ),
    "display_name": frozenset({"displayname", "nickname", "顯示名稱", "暱稱"}),
    "total_days": frozenset(
        {"count", "total", "totaldays", "checkincount", "簽到次數", "累積天數", "總天數"}
    ),
    "last_checkin_date": frozenset(
        {"lastdate", "lastcheckindate", "lastcheckin", "最後簽到日期", "最後日期"}
    ),
    "current_streak": frozenset({"streak", "currentstreak", "連續簽到", "連續天數"}),
    "daily_order": frozenset({"todayorder", "dailyorder", "今日順序", "當日順序"}),
}


def _normalize_header(value: object) -> str:
    text = str(value).strip().casefold()
    return re.sub(r"[\s_-]+", "", text)


_NORMALIZED_ALIASES = {
    _normalize_header(alias): canonical
    for canonical, aliases in _HEADER_ALIASES.items()
    for alias in aliases
}


def _safe_text(
    value: object,
    *,
    field: str,
    row_number: int,
    max_length: int = _MAX_CELL_LENGTH,
) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if len(text) > max_length:
        raise CheckinImportValidationError(f"第 {row_number} 列的 {field} 太長")
    return text


def _positive_int(value: object, *, field: str, row_number: int, allow_zero: bool) -> int:
    if isinstance(value, bool):
        raise CheckinImportValidationError(f"第 {row_number} 列的 {field} 必須是整數")
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float) and value.is_integer():
        parsed = int(value)
    else:
        text = _safe_text(value, field=field, row_number=row_number)
        if not re.fullmatch(r"[0-9]+", text):
            raise CheckinImportValidationError(f"第 {row_number} 列的 {field} 必須是整數")
        parsed = int(text)
    minimum = 0 if allow_zero else 1
    if parsed < minimum or parsed > 2_147_483_647:
        raise CheckinImportValidationError(f"第 {row_number} 列的 {field} 超出允許範圍")
    return parsed


def _iso_date(value: object, *, field: str, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _safe_text(value, field=field, row_number=row_number)
    try:
        parsed = date.fromisoformat(text)
    except ValueError as exc:
        raise CheckinImportValidationError(
            f"第 {row_number} 列的 {field} 必須是 YYYY-MM-DD"
        ) from exc
    if parsed.isoformat() != text:
        raise CheckinImportValidationError(f"第 {row_number} 列的 {field} 必須是 YYYY-MM-DD")
    return parsed


def _validated_headers(headers: list[object]) -> tuple[str, ...]:
    if not headers or all(not str(header).strip() for header in headers):
        raise CheckinImportValidationError("找不到欄位標題")
    if len(headers) > _MAX_COLUMNS:
        raise CheckinImportValidationError("欄位數超過上限")

    values: list[str] = []
    for raw in headers:
        text = "" if raw is None else str(raw).strip()
        if len(text) > _MAX_HEADER_LENGTH:
            raise CheckinImportValidationError("欄位標題太長")
        values.append(text)
    return tuple(values)


def _header_candidates(headers: tuple[str, ...]) -> dict[str, list[int]]:
    candidates: dict[str, list[int]] = {}
    for index, text in enumerate(headers):
        canonical = _NORMALIZED_ALIASES.get(_normalize_header(text))
        if canonical is None:
            continue
        candidates.setdefault(canonical, []).append(index)
    return candidates


def _suggested_mapping(headers: tuple[str, ...]) -> dict[str, int]:
    return {
        canonical: indexes[0]
        for canonical, indexes in _header_candidates(headers).items()
        if len(indexes) == 1
    }


def _map_headers(
    headers: list[object], column_mapping: dict[str, int] | None = None
) -> dict[str, int]:
    validated = _validated_headers(headers)
    mapped: dict[str, int]
    if column_mapping is None:
        candidates = _header_candidates(validated)
        ambiguous = next(
            (canonical for canonical, indexes in candidates.items() if len(indexes) > 1), None
        )
        if ambiguous is not None:
            raise CheckinImportValidationError(f"欄位 {ambiguous} 重複或意義不明")
        mapped = {canonical: indexes[0] for canonical, indexes in candidates.items()}
    else:
        mapped = {}
        used_indexes: set[int] = set()
        for canonical, index in column_mapping.items():
            if canonical not in _HEADER_ALIASES:
                raise CheckinImportValidationError("欄位對應包含未知的 Niibot 變數")
            if isinstance(index, bool) or not isinstance(index, int):
                raise CheckinImportValidationError("欄位對應位置必須是整數")
            if index < 0 or index >= len(validated) or not validated[index]:
                raise CheckinImportValidationError("欄位對應位置超出範圍")
            if index in used_indexes:
                raise CheckinImportValidationError("同一來源欄位不可對應多個 Niibot 變數")
            mapped[canonical] = index
            used_indexes.add(index)

    missing = {"total_days", "last_checkin_date"}.difference(mapped)
    if missing:
        raise CheckinImportValidationError("缺少 Count 或 LastDate 必要欄位")
    if not {"platform_user_id", "username", "display_name"}.intersection(mapped):
        raise CheckinImportValidationError("缺少觀眾識別欄位")
    return mapped


def _canonical_rows(
    rows: list[list[object]],
    *,
    through_date: date,
    column_mapping: dict[str, int] | None = None,
) -> tuple[tuple[SummaryRow | InvalidSummaryRow, ...], dict[str, int]]:
    if not rows:
        raise CheckinImportValidationError("匯入檔案是空的")
    mapping = _map_headers(rows[0], column_mapping)
    normalized: list[SummaryRow | InvalidSummaryRow] = []
    seen: set[str] = set()

    for row_number, row in enumerate(rows[1:], start=2):
        if all(value is None or not str(value).strip() for value in row):
            continue
        if len(normalized) >= _MAX_ROWS:
            raise CheckinImportValidationError("資料列數超過上限")

        def value(field: str, values: list[object] = row) -> object:
            index = mapping.get(field)
            return values[index] if index is not None and index < len(values) else None

        try:
            platform_user_id = (
                _safe_text(
                    value("platform_user_id"),
                    field="Twitch User ID",
                    row_number=row_number,
                )
                or None
            )
            username = (
                _safe_text(value("username"), field="Username", row_number=row_number) or None
            )
            display_name = (
                _safe_text(
                    value("display_name"),
                    field="DisplayName",
                    row_number=row_number,
                    max_length=128,
                )
                or None
            )
            if platform_user_id is not None and not platform_user_id.isascii():
                raise CheckinImportValidationError(f"第 {row_number} 列的 Twitch User ID 無效")
            if platform_user_id is not None and not platform_user_id.isdigit():
                raise CheckinImportValidationError(f"第 {row_number} 列的 Twitch User ID 無效")
            if username is not None and not _TWITCH_LOGIN.fullmatch(username):
                raise CheckinImportValidationError(f"第 {row_number} 列的 Username 無效")
            if platform_user_id is None and username is None and display_name is None:
                raise CheckinImportValidationError(f"第 {row_number} 列缺少觀眾身份")

            total_days = _positive_int(
                value("total_days"), field="Count", row_number=row_number, allow_zero=False
            )
            last_date = _iso_date(
                value("last_checkin_date"), field="LastDate", row_number=row_number
            )
            if last_date > through_date:
                raise CheckinImportValidationError(f"第 {row_number} 列的 LastDate 晚於截止日")

            streak_raw = value("current_streak")
            current_streak = None
            if streak_raw is not None and str(streak_raw).strip():
                current_streak = _positive_int(
                    streak_raw, field="Streak", row_number=row_number, allow_zero=True
                )
                if current_streak > total_days:
                    raise CheckinImportValidationError(f"第 {row_number} 列的 Streak 大於 Count")

            order_raw = value("daily_order")
            daily_order = None
            if order_raw is not None and str(order_raw).strip():
                daily_order = _positive_int(
                    order_raw, field="TodayOrder", row_number=row_number, allow_zero=False
                )

            if platform_user_id:
                key = f"id:{platform_user_id}"
            elif username:
                key = f"username:{username.casefold()}"
            else:
                key = f"display:{(display_name or '').casefold()}"
            if key in seen:
                raise CheckinImportValidationError(f"第 {row_number} 列的觀眾重複")
            seen.add(key)
            normalized.append(
                SummaryRow(
                    source_row=row_number,
                    platform_user_id=platform_user_id,
                    username=username,
                    display_name=display_name,
                    total_days=total_days,
                    last_checkin_date=last_date,
                    current_streak=current_streak,
                    daily_order=daily_order,
                )
            )
        except CheckinImportValidationError as exc:
            normalized.append(InvalidSummaryRow(source_row=row_number, issue=str(exc)))

    if not normalized:
        raise CheckinImportValidationError("沒有可匯入的資料列")
    return tuple(normalized), mapping


def _parse_text(content: bytes, *, delimiter: str) -> list[list[object]]:
    if len(content) > _MAX_TEXT_BYTES:
        raise CheckinImportValidationError("檔案大小超過 2 MiB 上限")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CheckinImportValidationError("文字檔必須使用 UTF-8 編碼") from exc
    try:
        rows: list[list[object]] = []
        for row_number, row in enumerate(
            csv.reader(io.StringIO(text), delimiter=delimiter, strict=True), start=1
        ):
            if row_number > _MAX_ROWS + 1:
                raise CheckinImportValidationError("資料列數超過上限")
            if len(row) > _MAX_COLUMNS:
                raise CheckinImportValidationError("欄位數超過上限")
            rows.append(list(row))
        return rows
    except csv.Error as exc:
        raise CheckinImportValidationError("文字表格格式無效") from exc


def _validate_xlsx_archive(content: bytes) -> None:
    if len(content) > _MAX_XLSX_BYTES:
        raise CheckinImportValidationError("XLSX 檔案大小超過 5 MiB 上限")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > _MAX_XLSX_ENTRIES:
                raise CheckinImportValidationError("XLSX 內含過多項目")
            total = 0
            for info in infos:
                path = PurePosixPath(info.filename)
                if path.is_absolute() or ".." in path.parts or info.flag_bits & 0x1:
                    raise CheckinImportValidationError("XLSX 封裝內容不安全")
                total += info.file_size
                if total > _MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise CheckinImportValidationError("XLSX 解壓後大小超過上限")
                if (info.file_size > 0 and info.compress_size == 0) or (
                    info.compress_size > 0
                    and info.file_size / info.compress_size > _MAX_XLSX_COMPRESSION_RATIO
                ):
                    raise CheckinImportValidationError("XLSX 壓縮比例超過上限")
    except BadZipFile as exc:
        raise CheckinImportValidationError("XLSX 檔案損壞") from exc


def _parse_xlsx(content: bytes, *, sheet_name: str | None) -> tuple[str, list[list[object]]]:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise CheckinImportValidationError("無法讀取 XLSX 檔案") from exc
    try:
        visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        if sheet_name:
            selected = next((sheet for sheet in visible if sheet.title == sheet_name), None)
            if selected is None:
                raise CheckinImportValidationError("找不到指定的工作表")
        else:
            selected = visible[0] if visible else None
            if selected is None:
                raise CheckinImportValidationError("XLSX 沒有可見的工作表")

        rows: list[list[object]] = []
        for row_index, cells in enumerate(selected.iter_rows(), start=1):
            if row_index > _MAX_ROWS + 1:
                raise CheckinImportValidationError("資料列數超過上限")
            if len(cells) > _MAX_COLUMNS:
                raise CheckinImportValidationError("欄位數超過上限")
            if any(cell.data_type == "f" for cell in cells):
                raise CheckinImportValidationError(f"第 {row_index} 列含公式，請先轉成值")
            rows.append([cell.value for cell in cells])
        return selected.title, rows
    finally:
        workbook.close()


def _tabular_rows(
    filename: str,
    content: bytes,
    *,
    sheet_name: str | None = None,
    source_format: Literal["google_sheets"] | None = None,
) -> tuple[Literal["csv", "tsv", "xlsx", "google_sheets"], str | None, list[list[object]]]:
    suffix = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    selected_sheet: str | None = None
    if source_format == "google_sheets":
        tabular_format: Literal["csv", "tsv", "xlsx", "google_sheets"] = "google_sheets"
        raw_rows = _parse_text(content, delimiter=",")
    elif suffix == "csv":
        tabular_format = "csv"
        raw_rows = _parse_text(content, delimiter=",")
    elif suffix == "tsv":
        tabular_format = "tsv"
        raw_rows = _parse_text(content, delimiter="\t")
    elif suffix == "xlsx":
        tabular_format = "xlsx"
        selected_sheet, raw_rows = _parse_xlsx(content, sheet_name=sheet_name)
    else:
        raise CheckinImportValidationError("只支援 CSV、TSV 或 XLSX 檔案")
    return tabular_format, selected_sheet, raw_rows


def inspect_summary_bytes(
    filename: str,
    content: bytes,
    *,
    sheet_name: str | None = None,
    source_format: Literal["google_sheets"] | None = None,
) -> InspectedSummary:
    """Return bounded source headers and exact alias suggestions without parsing viewer rows."""
    tabular_format, selected_sheet, raw_rows = _tabular_rows(
        filename,
        content,
        sheet_name=sheet_name,
        source_format=source_format,
    )
    if not raw_rows:
        raise CheckinImportValidationError("匯入檔案是空的")
    headers = _validated_headers(raw_rows[0])
    return InspectedSummary(
        format=tabular_format,
        sheet_name=selected_sheet,
        content_sha256=hashlib.sha256(content).hexdigest(),
        headers=headers,
        suggested_mapping=_suggested_mapping(headers),
    )


def parse_summary_bytes(
    filename: str,
    content: bytes,
    *,
    through_date: date,
    sheet_name: str | None = None,
    source_format: Literal["google_sheets"] | None = None,
    column_mapping: dict[str, int] | None = None,
) -> ParsedSummary:
    """Parse one upload or a fetched Google Sheets CSV into canonical rows."""
    tabular_format, selected_sheet, raw_rows = _tabular_rows(
        filename,
        content,
        sheet_name=sheet_name,
        source_format=source_format,
    )

    canonical_rows, resolved_mapping = _canonical_rows(
        raw_rows,
        through_date=through_date,
        column_mapping=column_mapping,
    )
    return ParsedSummary(
        format=tabular_format,
        sheet_name=selected_sheet,
        content_sha256=hashlib.sha256(content).hexdigest(),
        column_mapping=tuple(sorted(resolved_mapping.items())),
        rows=canonical_rows,
    )


def build_google_sheets_export_url(url: str) -> str:
    """Convert only a real Google Sheets document URL to a controlled CSV URL."""
    try:
        parsed = urlsplit(url)
    except ValueError as exc:
        raise CheckinImportValidationError("Google Sheets 連結無效") from exc
    try:
        port = parsed.port
    except ValueError as exc:
        raise CheckinImportValidationError("Google Sheets 連結無效") from exc
    if (
        parsed.scheme != "https"
        or parsed.hostname != _GOOGLE_HOST
        or port is not None
        or parsed.username is not None
        or parsed.password is not None
    ):
        raise CheckinImportValidationError("Google Sheets 連結無效")
    match = _GOOGLE_PATH.fullmatch(parsed.path)
    if match is None:
        raise CheckinImportValidationError("Google Sheets 連結無效")

    params = parse_qs(parsed.query, keep_blank_values=True)
    fragment_params = parse_qs(parsed.fragment, keep_blank_values=True)
    gids = set(params.get("gid", []) + fragment_params.get("gid", []))
    if len(gids) > 1 or (gids and not re.fullmatch(r"[0-9]{1,20}", next(iter(gids)))):
        raise CheckinImportValidationError("Google Sheets 工作表代碼無效")
    gid = next(iter(gids)) if gids else "0"
    path = f"/spreadsheets/d/{match.group(1)}/export"
    return urlunsplit(("https", _GOOGLE_HOST, path, urlencode({"format": "csv", "gid": gid}), ""))


def _google_export_redirect(location: str | None) -> str:
    if not location:
        raise CheckinImportValidationError("Google Sheets 匯出重新導向無效")
    try:
        parsed = urlsplit(location)
    except ValueError as exc:
        raise CheckinImportValidationError("Google Sheets 匯出重新導向無效") from exc
    try:
        hostname = (parsed.hostname or "").casefold()
        port = parsed.port
    except ValueError as exc:
        raise CheckinImportValidationError("Google Sheets 匯出重新導向無效") from exc
    if (
        parsed.scheme != "https"
        or port is not None
        or parsed.username is not None
        or parsed.password is not None
        or not _GOOGLE_EXPORT_HOST.fullmatch(hostname)
        or not parsed.path.startswith("/export/")
        or parsed.fragment
    ):
        raise CheckinImportValidationError("Google Sheets 匯出重新導向不安全")
    return location


async def _read_google_csv_response(response: httpx.Response) -> bytes:
    if response.status_code in {401, 403}:
        raise CheckinImportValidationError("這份 Google Sheets 尚未開放連結讀取")
    if response.status_code != 200:
        raise CheckinImportValidationError("Google Sheets 暫時無法讀取")
    content_type = response.headers.get("content-type", "").split(";", 1)[0].strip()
    if content_type not in {"text/csv", "text/plain", "application/octet-stream"}:
        raise CheckinImportValidationError("Google Sheets 回傳的格式不是 CSV")
    chunks: list[bytes] = []
    size = 0
    async for chunk in response.aiter_bytes():
        size += len(chunk)
        if size > _MAX_GOOGLE_BYTES:
            raise CheckinImportValidationError("Google Sheets 匯出大小超過上限")
        chunks.append(chunk)
    return b"".join(chunks)


async def fetch_google_sheet_csv(url: str, http: httpx.AsyncClient) -> bytes:
    """Fetch a bounded anonymous CSV export through one allowlisted Google redirect."""
    export_url = build_google_sheets_export_url(url)
    try:
        redirect_url: str | None = None
        async with http.stream(
            "GET", export_url, follow_redirects=False, timeout=httpx.Timeout(10.0)
        ) as response:
            if 300 <= response.status_code < 400:
                redirect_url = _google_export_redirect(response.headers.get("location"))
            else:
                return await _read_google_csv_response(response)

        async with http.stream(
            "GET", redirect_url, follow_redirects=False, timeout=httpx.Timeout(10.0)
        ) as response:
            if 300 <= response.status_code < 400:
                raise CheckinImportValidationError("Google Sheets 匯出發生多次重新導向")
            return await _read_google_csv_response(response)
    except CheckinImportValidationError:
        raise
    except httpx.HTTPError as exc:
        raise CheckinImportValidationError("Google Sheets 暫時無法讀取") from exc
